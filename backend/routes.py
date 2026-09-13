import base64
import csv
import io
import json
import secrets
import uuid
from datetime import datetime, timedelta, timezone

import psycopg2
import requests
from flask import jsonify, request, Response
from flask_bcrypt import Bcrypt

from . import eligible_source
from .auth import create_session, current_user
from .config import Config
from .constants import VALID_CAMPUSES, POPIA_CONSENT_VERSION
from .db import get_db, RealDictCursor
from .email_service import (
    send_verification_email,
    send_password_reset_email,
    send_booking_confirmation_email,
    send_verification_pending_email,
    send_verification_decision_email,
)

bcrypt = Bcrypt()

MAX_PDF_BYTES = 10 * 1024 * 1024
MAX_SPREADSHEET_BYTES = 5 * 1024 * 1024


# ============================================================
# Module-level helpers
# ============================================================

def _valid_campus(campus):
    return campus in VALID_CAMPUSES


def _call_verifynow_facematch(id_number, selfie_b64, reference_b64=None):
    if not Config.VERIFYNOW_API_KEY:
        return {"status": "error", "raw": "VERIFYNOW_API_KEY not configured"}

    body = {
        "mode": Config.VERIFYNOW_MODE,
        "idNumber": id_number,
        "selfieImageBase64": selfie_b64,
    }
    if reference_b64:
        body["referenceImageBase64"] = reference_b64

    try:
        r = requests.post(
            f"{Config.VERIFYNOW_API}/facematch",
            headers={
                "x-api-key": Config.VERIFYNOW_API_KEY,
                "Content-Type": "application/json",
                "Idempotency-Key": str(uuid.uuid4()),
            },
            json=body,
            timeout=30,
        )
        return r.json() if r.ok else {"status": "error", "raw": r.text}
    except Exception as e:
        return {"status": "error", "raw": str(e)}


def _log_admin_action(cur, admin, action, target_user_id=None,
                      target_student_number=None, target_full_name=None,
                      reason=None, metadata=None):
    cur.execute("""
        INSERT INTO admin_audit_log
            (admin_id, admin_student_number, action,
             target_user_id, target_student_number, target_full_name,
             reason, metadata)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
    """, (
        admin["id"], admin["student_number"], action,
        target_user_id, target_student_number, target_full_name,
        reason, json.dumps(metadata) if metadata else None,
    ))


# ---------- Spreadsheet import helpers ----------

def _norm_header(s):
    return "".join(ch for ch in str(s).lower() if ch.isalnum())


def _find_col(header, *aliases):
    wanted = {_norm_header(a) for a in aliases}
    for i, h in enumerate(header):
        if h in wanted:
            return i
    return None


def _find_col_contains(header, *fragments):
    frags = [_norm_header(f) for f in fragments]
    for i, h in enumerate(header):
        if all(f in h for f in frags):
            return i
    return None


def _clean_student_number(raw):
    s = str(raw).strip()
    if s.endswith(".0") and s[:-2].replace(".", "").isdigit():
        s = s[:-2]
    return s


CAMPUS_ALIASES = {
    "kwadlangenzwa":  "Kwadlangenzwa",
    "kwadlangezwa":   "Kwadlangenzwa",
    "kwa dlangezwa":  "Kwadlangenzwa",
    "kwa dlangenzwa": "Kwadlangenzwa",
    "richardsbay":    "Richards Bay",
    "richards bay":   "Richards Bay",
}


def _canonical_campus(raw):
    if not raw:
        return None
    key = " ".join(str(raw).split()).casefold()
    if key in CAMPUS_ALIASES:
        return CAMPUS_ALIASES[key]
    for c in VALID_CAMPUSES:
        if c.casefold() == key:
            return c
    return None


# ============================================================
# Routes
# ============================================================

def register_routes(app):
    bcrypt.init_app(app)

    # ---------- Eligibility ----------
    @app.post("/api/eligible-check")
    def eligible_check():
        data = request.get_json() or {}
        student_number = data.get("student_number", "").strip()
        if not student_number:
            return jsonify(error="Student number is required."), 400

        # 1) Google Sheet first
        sheet_student = eligible_source.get_sheet_student(student_number)
        if sheet_student:
            return jsonify(
                student_number=sheet_student["student_number"],
                campus=sheet_student["campus"],
                source="sheet",
            )

        # 2) Local DB fallback
        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                cur.execute("""
                    SELECT student_number, campus
                    FROM eligible_students WHERE student_number = %s
                """, (student_number,))
                row = cur.fetchone()

        if not row:
            return jsonify(error="This student number is not on the approved list."), 404

        return jsonify(
            student_number=row["student_number"],
            campus=row["campus"],
            source="database",
        )

    # ---------- Register ----------
    @app.post("/api/register")
    def register():
        data = request.get_json() or {}
        full_name = " ".join(data.get("full_name", "").strip().split())
        student_number = data.get("student_number", "").strip()
        email = data.get("email", "").strip().lower()
        password = data.get("password", "")
        campus = data.get("campus", "")
        popia_consent = bool(data.get("popia_consent", False))

        if not all([full_name, student_number, email, password, campus]):
            return jsonify(error="All registration fields are required."), 400
        if not _valid_campus(campus):
            return jsonify(error="Please select a valid campus."), 400
        if "@" not in email or "." not in email:
            return jsonify(error="Please enter a valid email address."), 400
        if len(password) < 8:
            return jsonify(error="Password must contain at least 8 characters."), 400
        if not popia_consent:
            return jsonify(error="You must agree to the POPIA Privacy Notice to create an account."), 400

        try:
            # 1) Canonical record: sheet first, DB fallback.
            canonical_name = None
            canonical_campus = None

            sheet_student = eligible_source.get_sheet_student(student_number)
            if sheet_student:
                canonical_name = sheet_student["full_name"]
                canonical_campus = sheet_student["campus"]
            else:
                with get_db() as conn:
                    with conn.cursor(cursor_factory=RealDictCursor) as cur:
                        cur.execute("""
                            SELECT full_name, campus FROM eligible_students
                            WHERE student_number = %s
                        """, (student_number,))
                        row = cur.fetchone()
                        if row:
                            canonical_name = row["full_name"]
                            canonical_campus = row["campus"]

            if not canonical_name:
                return jsonify(error="This student number is not on the approved student list."), 403

            if canonical_name.casefold() != full_name.casefold():
                return jsonify(error="The full name does not match the approved student record."), 403
            if canonical_campus != campus:
                return jsonify(error="The selected campus does not match the approved student record."), 403

            # 2) Upsert into eligible_students so the users FK is satisfied.
            with get_db() as conn:
                with conn.cursor(cursor_factory=RealDictCursor) as cur:
                    cur.execute("""
                        INSERT INTO eligible_students
                            (student_number, full_name, campus)
                        VALUES (%s, %s, %s)
                        ON CONFLICT (student_number) DO UPDATE
                            SET full_name = EXCLUDED.full_name,
                                campus = EXCLUDED.campus
                    """, (student_number, canonical_name, canonical_campus))

                    cur.execute("""
                        SELECT id FROM users WHERE student_number = %s OR email = %s
                    """, (student_number, email))
                    if cur.fetchone():
                        return jsonify(error="An account already exists for this student number or email."), 409

                    password_hash = bcrypt.generate_password_hash(password).decode("utf-8")
                    token = secrets.token_urlsafe(48)
                    expires = datetime.now(timezone.utc) + timedelta(hours=Config.TOKEN_HOURS)

                    cur.execute("""
                        INSERT INTO users
                            (student_number, full_name, email, password_hash, campus,
                             verification_token, verification_expires,
                             popia_consent, popia_consent_at, popia_consent_version)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, NOW(), %s)
                    """, (student_number, canonical_name, email,
                          password_hash, campus, token, expires,
                          True, POPIA_CONSENT_VERSION))
                conn.commit()

            try:
                send_verification_email(
                    email, canonical_name,
                    f"{Config.FRONTEND_URL}/verify.html?token={token}"
                )
                return jsonify(
                    message="Account created. Check your email to verify the account.",
                    needs_verification=True,
                ), 201
            except Exception:
                return jsonify(
                    message="Account created, but the verification email could not be sent.",
                    needs_verification=True,
                    warning="Check SMTP settings or request a new verification email.",
                ), 201

        except Exception:
            return jsonify(error="Database error while creating the account."), 500

    # ---------- Login ----------
    @app.post("/api/login")
    def login():
        data = request.get_json() or {}
        student_number = data.get("student_number", "").strip()
        password = data.get("password", "")

        if not student_number or not password:
            return jsonify(error="Student number and password are required."), 400

        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                cur.execute("""
                    SELECT id, student_number, full_name, email, password_hash,
                           campus, is_verified, popia_consent, stays_on_campus,
                           pathway, verification_status, is_admin, suspended
                    FROM users WHERE student_number = %s
                """, (student_number,))
                user = cur.fetchone()

                if not user or not bcrypt.check_password_hash(user["password_hash"], password):
                    return jsonify(error="Invalid student number or password."), 401

                if user["suspended"]:
                    return jsonify(
                        error="This account has been suspended. Contact Card Services."
                    ), 403

                if not user["is_verified"]:
                    return jsonify(
                        error="Please verify your email before logging in.",
                        needs_verification=True,
                    ), 403

                token = create_session(cur, user["id"])
            conn.commit()

        return jsonify(
            message="Login successful.",
            token=token,
            user={
                "student_number": user["student_number"],
                "full_name": user["full_name"],
                "email": user["email"],
                "campus": user["campus"],
                "popia_consent": user["popia_consent"],
                "stays_on_campus": user["stays_on_campus"],
                "pathway": user["pathway"],
                "verification_status": user["verification_status"],
                "is_admin": user["is_admin"],
            },
        )

    # ---------- Logout ----------
    @app.post("/api/logout")
    def logout():
        header = request.headers.get("Authorization", "")
        if not header.startswith("Bearer "):
            return jsonify(message="Already logged out."), 200
        token = header[7:].strip()

        with get_db() as conn:
            with conn.cursor() as cur:
                cur.execute("DELETE FROM sessions WHERE token = %s", (token,))
            conn.commit()

        return jsonify(message="Logged out."), 200

    # ---------- Current user ----------
    @app.get("/api/me")
    def me():
        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                user = current_user(cur)
                if not user:
                    return jsonify(error="Not signed in."), 401

        return jsonify(
            user={
                "student_number": user["student_number"],
                "full_name": user["full_name"],
                "email": user["email"],
                "campus": user["campus"],
                "popia_consent": user["popia_consent"],
                "stays_on_campus": user["stays_on_campus"],
                "pathway": user["pathway"],
                "verification_status": user["verification_status"],
                "is_admin": user["is_admin"],
                "created_at": user["created_at"].isoformat() if user["created_at"] else None,
            }
        )

    # ---------- Change password ----------
    @app.post("/api/change-password")
    def change_password():
        data = request.get_json() or {}
        current_password = data.get("current_password", "")
        new_password = data.get("new_password", "")

        if len(new_password) < 8:
            return jsonify(error="New password must contain at least 8 characters."), 400

        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                user = current_user(cur)
                if not user:
                    return jsonify(error="Not signed in."), 401

                cur.execute("SELECT password_hash FROM users WHERE id = %s", (user["id"],))
                row = cur.fetchone()
                if not row or not bcrypt.check_password_hash(row["password_hash"], current_password):
                    return jsonify(error="Current password is incorrect."), 401

                new_hash = bcrypt.generate_password_hash(new_password).decode("utf-8")
                cur.execute("UPDATE users SET password_hash = %s WHERE id = %s",
                            (new_hash, user["id"]))

                header = request.headers.get("Authorization", "")
                keep_token = header[7:].strip() if header.startswith("Bearer ") else ""
                cur.execute("""
                    DELETE FROM sessions WHERE user_id = %s AND token <> %s
                """, (user["id"], keep_token))
            conn.commit()

        return jsonify(message="Password changed successfully."), 200

    # ---------- Delete own account ----------
    @app.delete("/api/account")
    def delete_account():
        data = request.get_json() or {}
        password = data.get("password", "")

        if not password:
            return jsonify(error="Password is required to delete the account."), 400

        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                user = current_user(cur)
                if not user:
                    return jsonify(error="Not signed in."), 401

                cur.execute("SELECT password_hash FROM users WHERE id = %s", (user["id"],))
                row = cur.fetchone()
                if not row or not bcrypt.check_password_hash(row["password_hash"], password):
                    return jsonify(error="Password is incorrect."), 401

                cur.execute("DELETE FROM users WHERE id = %s", (user["id"],))
            conn.commit()

        return jsonify(message="Account deleted."), 200

    # ---------- Pathway choice ----------
    @app.post("/api/pathway")
    def set_pathway():
        data = request.get_json() or {}
        path = data.get("pathway")

        if path not in (1, 2):
            return jsonify(error="pathway must be 1 or 2."), 400

        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                user = current_user(cur)
                if not user:
                    return jsonify(error="Not signed in."), 401

                if path == 2:
                    cur.execute("""
                        UPDATE users
                        SET pathway = 2,
                            verification_status = 'not_required'
                        WHERE id = %s
                    """, (user["id"],))
                else:
                    cur.execute("UPDATE users SET pathway = 1 WHERE id = %s", (user["id"],))
            conn.commit()

        return jsonify(message="Pathway selected.", pathway=path), 200

    # ---------- Documents ----------
    @app.post("/api/documents/<doc_type>")
    def upload_document(doc_type):
        if doc_type not in ("id", "registration", "residence"):
            return jsonify(error="Unknown document type."), 400

        data = request.get_json() or {}
        filename = (data.get("filename") or "").strip()
        mime_type = (data.get("mime_type") or "").strip()
        content_b64 = data.get("content") or ""
        stays_on_campus = bool(data.get("stays_on_campus", False))

        if not filename or not content_b64:
            return jsonify(error="filename and content are required."), 400
        if not filename.lower().endswith(".pdf"):
            return jsonify(error="Only PDF files are accepted."), 400
        if mime_type and mime_type != "application/pdf":
            return jsonify(error="Only PDF files are accepted."), 400

        try:
            content_bytes = base64.b64decode(content_b64, validate=True)
        except Exception:
            return jsonify(error="File content is not valid base64."), 400

        if len(content_bytes) > MAX_PDF_BYTES:
            return jsonify(error="PDF is too large. Max 10 MB."), 413
        if not content_bytes.startswith(b"%PDF-"):
            return jsonify(error="File does not look like a valid PDF."), 400

        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                user = current_user(cur)
                if not user:
                    return jsonify(error="Not signed in."), 401

                pathway = user.get("pathway")

                if doc_type == "id" and pathway == 2:
                    return jsonify(
                        error="This pathway does not require an ID upload. "
                              "Staff will verify your ID on campus."
                    ), 400

                if doc_type == "residence" and not stays_on_campus:
                    return jsonify(
                        error="Proof of residence is only required for on-campus residents."
                    ), 400
                if doc_type == "residence" and stays_on_campus:
                    cur.execute(
                        "UPDATE users SET stays_on_campus = TRUE WHERE id = %s",
                        (user["id"],),
                    )

                cur.execute("""
                    INSERT INTO documents
                        (user_id, doc_type, filename, mime_type, size_bytes, content)
                    VALUES (%s, %s, %s, %s, %s, %s)
                    ON CONFLICT (user_id, doc_type) DO UPDATE
                        SET filename    = EXCLUDED.filename,
                            mime_type   = EXCLUDED.mime_type,
                            size_bytes  = EXCLUDED.size_bytes,
                            content     = EXCLUDED.content,
                            uploaded_at = NOW()
                """, (user["id"], doc_type, filename,
                      mime_type or "application/pdf",
                      len(content_bytes), psycopg2.Binary(content_bytes)))
            conn.commit()

        return jsonify(message="Document uploaded.", doc_type=doc_type), 201

    @app.get("/api/documents")
    def list_documents():
        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                user = current_user(cur)
                if not user:
                    return jsonify(error="Not signed in."), 401

                cur.execute("""
                    SELECT doc_type, filename, mime_type, size_bytes, uploaded_at
                    FROM documents WHERE user_id = %s
                """, (user["id"],))
                rows = cur.fetchall()

        return jsonify(documents=[
            {
                "doc_type": r["doc_type"],
                "filename": r["filename"],
                "mime_type": r["mime_type"],
                "size_bytes": r["size_bytes"],
                "uploaded_at": r["uploaded_at"].isoformat() if r["uploaded_at"] else None,
            }
            for r in rows
        ])

    @app.get("/api/documents/<doc_type>/download")
    def download_document(doc_type):
        if doc_type not in ("id", "registration", "residence"):
            return jsonify(error="Unknown document type."), 400

        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                user = current_user(cur)
                if not user:
                    return jsonify(error="Not signed in."), 401

                cur.execute("""
                    SELECT filename, mime_type, content
                    FROM documents WHERE user_id = %s AND doc_type = %s
                """, (user["id"], doc_type))
                row = cur.fetchone()

        if not row:
            return jsonify(error="Document not found."), 404

        return Response(
            bytes(row["content"]),
            mimetype=row["mime_type"] or "application/pdf",
            headers={"Content-Disposition": f'inline; filename="{row["filename"]}"'},
        )

    @app.delete("/api/documents/<doc_type>")
    def delete_document(doc_type):
        if doc_type not in ("id", "registration", "residence"):
            return jsonify(error="Unknown document type."), 400

        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                user = current_user(cur)
                if not user:
                    return jsonify(error="Not signed in."), 401

                cur.execute("""
                    DELETE FROM documents WHERE user_id = %s AND doc_type = %s
                """, (user["id"], doc_type))
            conn.commit()

        return jsonify(message="Document removed."), 200

    # ---------- Identity verification ----------
    @app.post("/api/verification/submit")
    def submit_verification():
        data = request.get_json() or {}
        id_number = (data.get("id_number") or "").strip()
        selfie_b64 = data.get("selfie_image") or ""
        id_photo_b64 = data.get("id_photo") or ""

        if not id_number or not selfie_b64:
            return jsonify(error="id_number and selfie_image are required."), 400

        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                user = current_user(cur)
                if not user:
                    return jsonify(error="Not signed in."), 401

                if user["pathway"] != 1:
                    return jsonify(
                        error="Identity verification is only required for the "
                              "online production pathway."
                    ), 400
                if user["verification_status"] == "verified":
                    return jsonify(error="Identity already verified."), 400

                result = _call_verifynow_facematch(
                    id_number, selfie_b64,
                    reference_b64=id_photo_b64 or None,
                )

                provider_status = (result.get("status") or "").lower()
                confidence = int(result.get("confidence_score") or 0)
                warnings = result.get("warnings") or []

                if provider_status == "match" and confidence >= Config.VERIFYNOW_AUTO_APPROVE_MIN:
                    outcome = "approved"
                elif "reference_photo_unavailable" in warnings or provider_status == "error":
                    outcome = "pending_review"
                elif confidence < Config.VERIFYNOW_MANUAL_REVIEW_MIN:
                    outcome = "declined"
                else:
                    outcome = "pending_review"

                cur.execute("""
                    INSERT INTO verifications
                        (user_id, provider, provider_txn_id, status,
                         confidence, warnings, raw_response)
                    VALUES (%s, 'verifynow', %s, %s, %s, %s, %s)
                    RETURNING id
                """, (
                    user["id"],
                    result.get("transaction_id") or result.get("reference"),
                    outcome,
                    confidence,
                    json.dumps(warnings),
                    json.dumps(result),
                ))
                verification_id = cur.fetchone()["id"]

                user_status = {
                    "approved": "verified",
                    "declined": "rejected",
                    "pending_review": "pending_review",
                }[outcome]

                cur.execute("UPDATE users SET verification_status = %s WHERE id = %s",
                            (user_status, user["id"]))
            conn.commit()

        if outcome == "pending_review":
            send_verification_pending_email(user["email"], user["full_name"])

        return jsonify(
            verification_id=verification_id,
            outcome=outcome,
            confidence=confidence,
            warnings=warnings,
        ), 200

    @app.get("/api/verification/me")
    def my_verification():
        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                user = current_user(cur)
                if not user:
                    return jsonify(error="Not signed in."), 401

                cur.execute("""
                    SELECT id, status, confidence, warnings,
                           admin_decision, admin_notes, created_at, reviewed_at
                    FROM verifications
                    WHERE user_id = %s
                    ORDER BY created_at DESC LIMIT 1
                """, (user["id"],))
                row = cur.fetchone()

        if not row:
            return jsonify(verification=None)

        return jsonify(verification={
            "id": row["id"],
            "status": row["status"],
            "confidence": row["confidence"],
            "warnings": row["warnings"],
            "admin_decision": row["admin_decision"],
            "admin_notes": row["admin_notes"],
            "created_at": row["created_at"].isoformat() if row["created_at"] else None,
            "reviewed_at": row["reviewed_at"].isoformat() if row["reviewed_at"] else None,
        })

    # ---------- Admin: verification queue ----------
    @app.get("/api/admin/verifications")
    def admin_list_verifications():
        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                user = current_user(cur)
                if not user or not user.get("is_admin"):
                    return jsonify(error="Admin access required."), 403

                cur.execute("""
                    SELECT v.id, v.status, v.confidence, v.warnings, v.created_at,
                           u.id AS user_id, u.student_number, u.full_name,
                           u.email, u.campus
                    FROM verifications v
                    JOIN users u ON u.id = v.user_id
                    WHERE v.status = 'pending_review'
                    ORDER BY v.created_at ASC
                """)
                rows = cur.fetchall()

        return jsonify(pending=[
            {
                "id": r["id"],
                "user_id": r["user_id"],
                "student_number": r["student_number"],
                "full_name": r["full_name"],
                "email": r["email"],
                "campus": r["campus"],
                "confidence": r["confidence"],
                "warnings": r["warnings"],
                "created_at": r["created_at"].isoformat(),
            }
            for r in rows
        ])

    @app.post("/api/admin/verifications/<int:vid>/decide")
    def admin_decide_verification(vid):
        data = request.get_json() or {}
        decision = (data.get("decision") or "").lower()
        notes = (data.get("notes") or "").strip()

        if decision not in ("approve", "decline"):
            return jsonify(error="decision must be 'approve' or 'decline'."), 400

        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                admin = current_user(cur)
                if not admin or not admin.get("is_admin"):
                    return jsonify(error="Admin access required."), 403

                cur.execute("SELECT user_id FROM verifications WHERE id = %s", (vid,))
                row = cur.fetchone()
                if not row:
                    return jsonify(error="Verification not found."), 404

                target_user_id = row["user_id"]

                cur.execute("""
                    UPDATE verifications
                    SET status = %s,
                        reviewed_by = %s,
                        reviewed_at = NOW(),
                        admin_decision = %s,
                        admin_notes = %s
                    WHERE id = %s
                """, (
                    "approved" if decision == "approve" else "declined",
                    admin["id"],
                    decision,
                    notes or None,
                    vid,
                ))
                cur.execute("UPDATE users SET verification_status = %s WHERE id = %s",
                            ("verified" if decision == "approve" else "rejected",
                             target_user_id))

                cur.execute("SELECT student_number, full_name, email FROM users WHERE id = %s",
                            (target_user_id,))
                target = cur.fetchone()

                _log_admin_action(
                    cur, admin,
                    "approve" if decision == "approve" else "decline",
                    target_user_id=target_user_id,
                    target_student_number=target["student_number"],
                    target_full_name=target["full_name"],
                    reason=notes or None,
                )
            conn.commit()

        if target:
            send_verification_decision_email(
                target["email"], target["full_name"],
                approved=(decision == "approve"),
                notes=notes,
            )

        return jsonify(message=f"Verification {decision}d.", verification_id=vid), 200

    # ---------- Admin: list users ----------
    @app.get("/api/admin/users")
    def admin_list_users():
        q = (request.args.get("q") or "").strip()
        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                admin = current_user(cur)
                if not admin or not admin.get("is_admin"):
                    return jsonify(error="Admin access required."), 403

                if q:
                    like = f"%{q}%"
                    cur.execute("""
                        SELECT id, student_number, full_name, email, campus,
                               pathway, verification_status, is_admin, suspended,
                               created_at
                        FROM users
                        WHERE student_number ILIKE %s
                           OR full_name ILIKE %s
                           OR email ILIKE %s
                        ORDER BY created_at DESC
                        LIMIT 200
                    """, (like, like, like))
                else:
                    cur.execute("""
                        SELECT id, student_number, full_name, email, campus,
                               pathway, verification_status, is_admin, suspended,
                               created_at
                        FROM users
                        ORDER BY created_at DESC
                        LIMIT 200
                    """)
                rows = cur.fetchall()

        return jsonify(users=[
            {
                "id": r["id"],
                "student_number": r["student_number"],
                "full_name": r["full_name"],
                "email": r["email"],
                "campus": r["campus"],
                "pathway": r["pathway"],
                "verification_status": r["verification_status"],
                "is_admin": r["is_admin"],
                "suspended": r["suspended"],
                "created_at": r["created_at"].isoformat() if r["created_at"] else None,
            }
            for r in rows
        ])

    # ---------- Admin: suspend / reactivate ----------
    @app.post("/api/admin/users/<int:uid>/suspend")
    def admin_suspend_user(uid):
        data = request.get_json() or {}
        suspend = bool(data.get("suspend", True))
        reason = (data.get("reason") or "").strip()

        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                admin = current_user(cur)
                if not admin or not admin.get("is_admin"):
                    return jsonify(error="Admin access required."), 403
                if admin["id"] == uid:
                    return jsonify(error="You cannot suspend your own account."), 400

                cur.execute("""
                    SELECT student_number, full_name FROM users WHERE id = %s
                """, (uid,))
                target = cur.fetchone()
                if not target:
                    return jsonify(error="User not found."), 404

                cur.execute("UPDATE users SET suspended = %s WHERE id = %s",
                            (suspend, uid))

                if suspend:
                    cur.execute("DELETE FROM sessions WHERE user_id = %s", (uid,))

                _log_admin_action(
                    cur, admin,
                    "suspend" if suspend else "reactivate",
                    target_user_id=uid,
                    target_student_number=target["student_number"],
                    target_full_name=target["full_name"],
                    reason=reason or None,
                )
            conn.commit()

        return jsonify(message=("User suspended." if suspend else "User reactivated.")), 200

    # ---------- Admin: delete account ----------
    @app.delete("/api/admin/users/<int:uid>")
    def admin_delete_user(uid):
        data = request.get_json() or {}
        reason = (data.get("reason") or "").strip()

        if not reason:
            return jsonify(error="A reason is required when deleting an account."), 400

        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                admin = current_user(cur)
                if not admin or not admin.get("is_admin"):
                    return jsonify(error="Admin access required."), 403
                if admin["id"] == uid:
                    return jsonify(error="You cannot delete your own account from here."), 400

                cur.execute("""
                    SELECT student_number, full_name FROM users WHERE id = %s
                """, (uid,))
                target = cur.fetchone()
                if not target:
                    return jsonify(error="User not found."), 404

                _log_admin_action(
                    cur, admin, "delete",
                    target_user_id=uid,
                    target_student_number=target["student_number"],
                    target_full_name=target["full_name"],
                    reason=reason,
                )

                cur.execute("DELETE FROM users WHERE id = %s", (uid,))
            conn.commit()

        return jsonify(message="Account deleted."), 200

    # ---------- Admin: bulk eligible upload ----------
    @app.post("/api/admin/eligible/upload")
    def admin_upload_eligible():
        from openpyxl import load_workbook

        data = request.get_json() or {}
        filename = (data.get("filename") or "").strip()
        content_b64 = data.get("content") or ""

        if not filename or not content_b64:
            return jsonify(error="filename and content are required."), 400

        try:
            raw = base64.b64decode(content_b64, validate=True)
        except Exception:
            return jsonify(error="File content is not valid base64."), 400

        if len(raw) > MAX_SPREADSHEET_BYTES:
            return jsonify(error="File is too large. Max 5 MB."), 413

        ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

        rows = []
        try:
            if ext == "csv":
                text = raw.decode("utf-8-sig", errors="replace")
                rows = list(csv.reader(io.StringIO(text)))
            elif ext in ("xlsx", "xlsm"):
                wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
                sheet = wb.active
                for r in sheet.iter_rows(values_only=True):
                    rows.append([("" if c is None else c) for c in r])
            else:
                return jsonify(error="Only .xlsx, .xlsm, or .csv files are accepted."), 400
        except Exception as e:
            return jsonify(error=f"Could not read file: {e}"), 400

        if not rows:
            return jsonify(error="The file is empty."), 400

        header = [_norm_header(c) for c in rows[0]]

        col_studno = _find_col(
            header,
            "student_number", "student number", "studentno", "studentnumber",
        )
        if col_studno is None:
            col_studno = _find_col_contains(header, "student", "number")

        col_name = _find_col(header, "full_name", "full name", "name")
        col_first = _find_col(header, "first name", "firstname", "first names", "given name")
        col_middle = _find_col(header, "middle names", "middlenames", "middle name", "middle")
        col_last = _find_col(header, "surnames", "surname", "last name", "lastname", "family name")

        col_campus = _find_col(header, "campus", "campus_name", "campusname")
        if col_campus is None:
            col_campus = _find_col_contains(header, "campus")

        if col_studno is None:
            return jsonify(error="Could not find a student number column in the file."), 400
        if col_name is None and col_first is None and col_last is None:
            return jsonify(error="Could not find a name column in the file."), 400
        if col_campus is None:
            return jsonify(error="Could not find a campus column in the file."), 400

        def build_full_name(row):
            if col_name is not None and col_name < len(row) and row[col_name]:
                return " ".join(str(row[col_name]).split())
            parts = []
            for c in (col_first, col_middle, col_last):
                if c is not None and c < len(row) and row[c]:
                    parts.append(str(row[c]).strip())
            return " ".join(" ".join(parts).split())

        to_upsert = {}
        skipped = []

        for idx, row in enumerate(rows[1:], start=2):
            if row is None:
                continue

            def cell(i):
                return str(row[i]).strip() if i < len(row) and row[i] is not None else ""

            studno = _clean_student_number(cell(col_studno))
            name = build_full_name(row)
            campus = cell(col_campus)

            if not studno and not name and not campus:
                continue

            if not studno:
                skipped.append({"row": idx, "reason": "Missing student number"}); continue
            if not name:
                skipped.append({"row": idx, "reason": "Missing full name"}); continue
            if not campus:
                skipped.append({"row": idx, "reason": "Missing campus"}); continue

            campus_value = _canonical_campus(campus)
            if not campus_value:
                skipped.append({
                    "row": idx,
                    "reason": f"Campus must be 'Kwadlangenzwa' or 'Richards Bay' (got '{campus}')"
                })
                continue

            to_upsert[studno] = (studno, name, campus_value)

        if not to_upsert:
            return jsonify(
                error="No valid rows to import.",
                skipped=skipped[:50],
                skipped_count=len(skipped),
            ), 400

        inserted = 0
        updated = 0
        unchanged = 0

        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                admin = current_user(cur)
                if not admin or not admin.get("is_admin"):
                    return jsonify(error="Admin access required."), 403

                for studno, (num, name, campus) in to_upsert.items():
                    cur.execute("""
                        SELECT full_name, campus FROM eligible_students
                        WHERE student_number = %s
                    """, (num,))
                    existing = cur.fetchone()

                    if existing is None:
                        cur.execute("""
                            INSERT INTO eligible_students
                                (student_number, full_name, campus)
                            VALUES (%s, %s, %s)
                        """, (num, name, campus))
                        inserted += 1
                    elif existing["full_name"] != name or existing["campus"] != campus:
                        cur.execute("""
                            UPDATE eligible_students
                            SET full_name = %s, campus = %s
                            WHERE student_number = %s
                        """, (name, campus, num))
                        updated += 1
                    else:
                        unchanged += 1

                _log_admin_action(
                    cur, admin, "eligible_upload",
                    reason=f"Uploaded {filename}",
                    metadata={
                        "filename": filename,
                        "inserted": inserted,
                        "updated": updated,
                        "unchanged": unchanged,
                        "skipped": len(skipped),
                    },
                )
            conn.commit()

        return jsonify(
            message="Upload complete.",
            filename=filename,
            total_valid=len(to_upsert),
            inserted=inserted,
            updated=updated,
            unchanged=unchanged,
            skipped=skipped[:50],
            skipped_count=len(skipped),
        ), 200

    # ---------- Admin: list eligible ----------
    @app.get("/api/admin/eligible")
    def admin_list_eligible():
        q = (request.args.get("q") or "").strip()
        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                admin = current_user(cur)
                if not admin or not admin.get("is_admin"):
                    return jsonify(error="Admin access required."), 403

                if q:
                    like = f"%{q}%"
                    cur.execute("""
                        SELECT student_number, full_name, campus
                        FROM eligible_students
                        WHERE student_number ILIKE %s
                           OR full_name ILIKE %s
                        ORDER BY student_number ASC
                        LIMIT 500
                    """, (like, like))
                else:
                    cur.execute("""
                        SELECT student_number, full_name, campus
                        FROM eligible_students
                        ORDER BY student_number ASC
                        LIMIT 500
                    """)
                rows = cur.fetchall()

        return jsonify(eligible=[dict(r) for r in rows])

    # ---------- Admin: delete eligible ----------
    @app.delete("/api/admin/eligible/<student_number>")
    def admin_delete_eligible(student_number):
        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                admin = current_user(cur)
                if not admin or not admin.get("is_admin"):
                    return jsonify(error="Admin access required."), 403

                cur.execute("""
                    SELECT id FROM users WHERE student_number = %s
                """, (student_number,))
                if cur.fetchone():
                    return jsonify(
                        error="This student already has an account. "
                              "Suspend or delete the account first."
                    ), 409

                cur.execute("""
                    DELETE FROM eligible_students WHERE student_number = %s
                """, (student_number,))
                deleted = cur.rowcount

                if deleted:
                    _log_admin_action(
                        cur, admin, "eligible_delete",
                        target_student_number=student_number,
                        reason="Removed from approved list",
                    )
            conn.commit()

        if not deleted:
            return jsonify(error="Student number not found."), 404

        return jsonify(message="Eligible student removed."), 200

    # ---------- Admin: eligibility source status ----------
    @app.get("/api/admin/eligible/source")
    def admin_eligible_source_status():
        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                admin = current_user(cur)
                if not admin or not admin.get("is_admin"):
                    return jsonify(error="Admin access required."), 403

        return jsonify(sheet=eligible_source.cache_metadata())

    # ---------- Admin: sync from Google Sheet ----------
    @app.post("/api/admin/eligible/sync")
    def admin_eligible_sync():
        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                admin = current_user(cur)
                if not admin or not admin.get("is_admin"):
                    return jsonify(error="Admin access required."), 403

                eligible_source.refresh_cache(force=True)
                students = eligible_source.cached_students()

                if not students:
                    meta = eligible_source.cache_metadata()
                    return jsonify(
                        error="No students found in the Google Sheet.",
                        sheet=meta,
                    ), 400

                inserted = 0
                updated = 0
                unchanged = 0

                for studno, s in students.items():
                    cur.execute("""
                        SELECT full_name, campus FROM eligible_students
                        WHERE student_number = %s
                    """, (studno,))
                    existing = cur.fetchone()

                    if existing is None:
                        cur.execute("""
                            INSERT INTO eligible_students
                                (student_number, full_name, campus)
                            VALUES (%s, %s, %s)
                        """, (studno, s["full_name"], s["campus"]))
                        inserted += 1
                    elif existing["full_name"] != s["full_name"] or existing["campus"] != s["campus"]:
                        cur.execute("""
                            UPDATE eligible_students
                            SET full_name = %s, campus = %s
                            WHERE student_number = %s
                        """, (s["full_name"], s["campus"], studno))
                        updated += 1
                    else:
                        unchanged += 1

                _log_admin_action(
                    cur, admin, "eligible_sync",
                    reason="Synced from Google Sheet",
                    metadata={
                        "inserted": inserted,
                        "updated": updated,
                        "unchanged": unchanged,
                    },
                )
            conn.commit()

        return jsonify(
            message="Synced from Google Sheet.",
            inserted=inserted,
            updated=updated,
            unchanged=unchanged,
            total=len(students),
        ), 200

    # ---------- Admin: audit log ----------
    @app.get("/api/admin/audit")
    def admin_audit_log():
        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                admin = current_user(cur)
                if not admin or not admin.get("is_admin"):
                    return jsonify(error="Admin access required."), 403

                cur.execute("""
                    SELECT id, admin_student_number, action,
                           target_user_id, target_student_number, target_full_name,
                           reason, created_at
                    FROM admin_audit_log
                    ORDER BY created_at DESC
                    LIMIT 200
                """)
                rows = cur.fetchall()

        return jsonify(entries=[
            {
                "id": r["id"],
                "admin_student_number": r["admin_student_number"],
                "action": r["action"],
                "target_user_id": r["target_user_id"],
                "target_student_number": r["target_student_number"],
                "target_full_name": r["target_full_name"],
                "reason": r["reason"],
                "created_at": r["created_at"].isoformat() if r["created_at"] else None,
            }
            for r in rows
        ])

    # ---------- Verify email ----------
    @app.get("/api/verify")
    def verify():
        token = request.args.get("token", "")
        if not token:
            return jsonify(error="Missing verification token."), 400

        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                cur.execute("""
                    SELECT id FROM users
                    WHERE verification_token = %s
                      AND verification_expires > NOW()
                """, (token,))
                user = cur.fetchone()

                if not user:
                    return jsonify(error="Invalid or expired verification link."), 400

                cur.execute("""
                    UPDATE users
                    SET is_verified = TRUE,
                        verification_token = NULL,
                        verification_expires = NULL
                    WHERE id = %s
                """, (user["id"],))
            conn.commit()

        return jsonify(message="Account verified successfully. You can now log in.")

    # ---------- Resend verification ----------
    @app.post("/api/resend-verification")
    def resend_verification():
        data = request.get_json() or {}
        student_number = data.get("student_number", "").strip()

        if not student_number:
            return jsonify(error="Student number is required."), 400

        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                cur.execute("""
                    SELECT id, full_name, email, is_verified
                    FROM users WHERE student_number = %s
                """, (student_number,))
                user = cur.fetchone()

                if not user:
                    return jsonify(message="If an account exists and is unverified, a new link has been sent."), 200
                if user["is_verified"]:
                    return jsonify(error="This account is already verified. You can log in."), 400

                token = secrets.token_urlsafe(48)
                expires = datetime.now(timezone.utc) + timedelta(hours=Config.TOKEN_HOURS)

                cur.execute("""
                    UPDATE users
                    SET verification_token = %s, verification_expires = %s
                    WHERE id = %s
                """, (token, expires, user["id"]))
            conn.commit()

        try:
            send_verification_email(
                user["email"], user["full_name"],
                f"{Config.FRONTEND_URL}/verify.html?token={token}",
            )
            return jsonify(message="A new verification link has been sent to your email."), 200
        except Exception:
            return jsonify(error="Could not send the verification email. Check SMTP settings."), 500

    # ---------- Booking confirmation email ----------
    @app.post("/api/send-booking-confirmation")
    def send_booking_confirmation():
        data = request.get_json() or {}
        student_number = data.get("student_number", "").strip()

        if not student_number:
            return jsonify(error="Student number is required."), 400

        required = ["service_label", "date_str", "time_slot", "queue_no", "ref_code", "pathway"]
        if not all(data.get(f) is not None for f in required):
            return jsonify(error="Missing booking details."), 400

        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                cur.execute("""
                    SELECT full_name, email, campus
                    FROM users WHERE student_number = %s
                """, (student_number,))
                user = cur.fetchone()

                if not user:
                    return jsonify(error="Student not found."), 404

                cur.execute("""
                    UPDATE users
                    SET popia_consent = TRUE,
                        popia_consent_at = NOW(),
                        popia_consent_version = %s
                    WHERE student_number = %s
                """, (POPIA_CONSENT_VERSION, student_number))
            conn.commit()

        try:
            send_booking_confirmation_email(
                user["email"],
                user["full_name"],
                data["service_label"],
                user["campus"],
                data["date_str"],
                data["time_slot"],
                data["queue_no"],
                data["ref_code"],
                int(data["pathway"]),
            )
            return jsonify(message="Booking confirmation email sent."), 200
        except Exception:
            return jsonify(error="Could not send the confirmation email."), 500

    # ---------- Forgot / reset password ----------
    @app.post("/api/forgot-password")
    def forgot_password():
        data = request.get_json() or {}
        email = data.get("email", "").strip().lower()

        if not email:
            return jsonify(error="Email address is required."), 400

        with get_db() as conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                cur.execute("SELECT id, full_name FROM users WHERE email = %s", (email,))
                user = cur.fetchone()

                if user:
                    token = secrets.token_urlsafe(48)
                    expires = datetime.now(timezone.utc) + timedelta(hours=1)

                    cur.execute("""
                        UPDATE users SET reset_token = %s, reset_expires = %s
                        WHERE id = %s
                    """, (token, expires, user["id"]))
                    conn.commit()

                    send_password_reset_email(
                        email, user["full_name"],
                        f"{Config.FRONTEND_URL}/reset.html?token={token}"
                    )

        return jsonify(message="If an account exists for that email, a password reset link has been sent.")

    @app.post("/api/reset-password")
    def reset_password():
        data = request.get_json() or {}
        token = data.get("token", "")
        new_password = data.get("password", "")

        if not token or len(new_password) < 8:
            return jsonify(error="A valid token and password of at least 8 characters are required."), 400

        password_hash = bcrypt.generate_password_hash(new_password).decode("utf-8")

        with get_db() as conn:
            with conn.cursor() as cur:
                cur.execute("""
                    UPDATE users
                    SET password_hash = %s,
                        reset_token = NULL,
                        reset_expires = NULL
                    WHERE reset_token = %s
                      AND reset_expires > NOW()
                """, (password_hash, token))
                changed = cur.rowcount
            conn.commit()

        if changed != 1:
            return jsonify(error="Invalid or expired password reset link."), 400

        return jsonify(message="Password changed successfully. You can now log in.")